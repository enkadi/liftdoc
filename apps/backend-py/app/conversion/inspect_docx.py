"""Read a generated DOCX back to measure what was actually written.

Word counts, reconstructed-list counts and the page estimate are taken from the output file
itself rather than from the intermediate model, so the fidelity report describes the artefact
the user downloads.
"""

from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path

import pymupdf
from docx import Document
from docx.oxml.ns import qn

from .docx_writer import SINGLE_LINE_FACTOR
from .extract import count_words

_W_P, _W_TBL, _W_T, _W_TAB, _W_BR = (qn(f"w:{t}") for t in ("p", "tbl", "t", "tab", "br"))
EMU_PER_PT = 12700
CELL_PADDING = 11.0  # default left+right cell margins, points


@dataclass(slots=True)
class DocxStats:
    word_count: int
    page_count: int
    ordered_lists: int
    unordered_lists: int
    tables: int
    images: int
    headings: dict[int, int]
    explicit_pages: int = 1  # pages delimited by explicit page/section breaks


def _paragraph_text(p) -> str:
    parts: list[str] = []
    for node in p.iter(_W_T, _W_TAB, _W_BR):
        parts.append(node.text or "" if node.tag == _W_T else " ")
    return "".join(parts)


class _Styles:
    def __init__(self, doc) -> None:
        self._doc = doc
        self._cache: dict[str, tuple[float, float, float, float]] = {}

    def metrics(self, style) -> tuple[float, float, float, float]:
        """(font size, line spacing multiple, space before, space after) via the style chain."""
        if style.name in self._cache:
            return self._cache[style.name]
        size = spacing = before = after = None
        s = style
        while s is not None:
            if size is None and s.font.size is not None:
                size = s.font.size.pt
            pf = s.paragraph_format
            if spacing is None and pf.line_spacing is not None:
                spacing = pf.line_spacing
            if before is None and pf.space_before is not None:
                before = pf.space_before.pt
            if after is None and pf.space_after is not None:
                after = pf.space_after.pt
            s = s.base_style
        # a Length line spacing ("at least"/"exactly") is kept negative to tell it from a multiple
        if spacing is not None and not isinstance(spacing, float):
            spacing = -spacing.pt
        result = (size or 11.0, spacing if spacing is not None else 1.0, before or 0.0, after or 0.0)
        self._cache[style.name] = result
        return result


def _text_width(text: str, size: float) -> float:
    try:
        return pymupdf.get_text_length(text, fontname="helv", fontsize=size)
    except Exception:  # noqa: BLE001
        return len(text) * size * 0.5


def inspect_docx(path: Path) -> DocxStats:
    doc = Document(str(path))
    body = doc.element.body
    styles = _Styles(doc)

    # ---- numbering: which numIds are ordered at their top level
    ordered_by_num: dict[str, bool] = {}
    try:
        numbering = doc.part.numbering_part.element
        fmt_by_abstract = {}
        for a in numbering.findall(qn("w:abstractNum")):
            lvl0 = a.find(qn("w:lvl"))
            fmt = lvl0.find(qn("w:numFmt")).get(qn("w:val")) if lvl0 is not None else "bullet"
            fmt_by_abstract[a.get(qn("w:abstractNumId"))] = fmt
        fmt_by_level: dict[tuple[str, str], str] = {}
        for a in numbering.findall(qn("w:abstractNum")):
            for lvl in a.findall(qn("w:lvl")):
                fmt_by_level[(a.get(qn("w:abstractNumId")), lvl.get(qn("w:ilvl")))] = lvl.find(
                    qn("w:numFmt")
                ).get(qn("w:val"))
        abstract_of = {
            n.get(qn("w:numId")): n.find(qn("w:abstractNumId")).get(qn("w:val"))
            for n in numbering.findall(qn("w:num"))
        }
    except Exception:  # noqa: BLE001 - a document without numbering has no lists
        fmt_by_abstract, fmt_by_level, abstract_of = {}, {}, {}

    words = 0
    used_nums: set[str] = set()
    headings: dict[int, int] = {}
    for p in body.iter(_W_P):
        words += count_words(_paragraph_text(p).split())
        num_pr = p.find(f"{qn('w:pPr')}/{qn('w:numPr')}")
        if num_pr is not None:
            num_id = num_pr.find(qn("w:numId")).get(qn("w:val"))
            ilvl_el = num_pr.find(qn("w:ilvl"))
            ilvl = ilvl_el.get(qn("w:val")) if ilvl_el is not None else "0"
            used_nums.add(num_id)
            if fmt_by_level.get((abstract_of.get(num_id), ilvl), "bullet") != "bullet":
                words += 1  # the rendered number ("1.") is visible text in the output
        style = p.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
        if style is not None:
            val = style.get(qn("w:val"), "")
            if val.startswith("Heading") and val[-1:].isdigit():
                headings[int(val[-1])] = headings.get(int(val[-1]), 0) + 1
    for num_id in used_nums:
        ordered_by_num[num_id] = fmt_by_abstract.get(abstract_of.get(num_id), "bullet") != "bullet"

    first = doc.sections[0]
    for part in (first.header, first.footer):
        for p in part.paragraphs:
            words += count_words(_paragraph_text(p._p).split())

    return DocxStats(
        word_count=words,
        page_count=_estimate_pages(doc, styles),
        ordered_lists=sum(1 for v in ordered_by_num.values() if v),
        unordered_lists=sum(1 for v in ordered_by_num.values() if not v),
        tables=len(body.findall(_W_TBL)),
        images=len(list(body.iter(qn("wp:inline"), qn("wp:anchor")))),
        headings=headings,
        explicit_pages=_explicit_pages(body),
    )


def _explicit_pages(body) -> int:
    pages = 1
    for p in body.iter(_W_P):
        ppr = p.find(qn("w:pPr"))
        if ppr is not None and ppr.find(qn("w:pageBreakBefore")) is not None:
            pages += 1
        sect = ppr.find(qn("w:sectPr")) if ppr is not None else None
        if sect is not None:
            # the type of a break is stored on the section that FOLLOWS it; a paragraph-level
            # sectPr closes a section, so look at the next one
            following = None
            for later in p.itersiblings():
                found = later.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") if later.tag == _W_P else None
                if found is not None or later.tag == qn("w:sectPr"):
                    following = found if found is not None else later
                    break
            kind = following.find(qn("w:type")) if following is not None else None
            if kind is None or kind.get(qn("w:val")) != "continuous":
                pages += 1
    pages += len(body.findall(f".//{qn('w:br')}[@{qn('w:type')}='page']"))
    return pages


def _estimate_pages(doc, styles: _Styles) -> int:
    """Best-effort pagination: lay blocks out on pages using approximate text metrics."""
    sections = [
        (
            (s.page_width.pt - s.left_margin.pt - s.right_margin.pt),
            (s.page_height.pt - s.top_margin.pt - s.bottom_margin.pt),
        )
        for s in doc.sections
    ]
    # (column count, starts on the same page) per section, in document order
    layouts = []
    for s in doc.sections:
        cols = s._sectPr.find(qn("w:cols"))
        kind = s._sectPr.find(qn("w:type"))
        layouts.append(
            (
                max(1, int(cols.get(qn("w:num"), 1))) if cols is not None else 1,
                kind is not None and kind.get(qn("w:val")) == "continuous",
            )
        )
    pages, y, sec = 1, 0.0, 0
    columns = layouts[0][0]
    width, height = sections[0]
    normal = doc.styles["Normal"]

    def place(block_height: float, line_height: float = 0.0) -> None:
        nonlocal pages, y
        remaining = height - y
        if block_height <= remaining + 0.5:
            y += block_height
            return
        if line_height and block_height > 2 * line_height:
            fit = int(remaining // line_height)
            rest = block_height - (fit * line_height if fit >= 2 else 0.0)
        else:
            rest = block_height
        pages += 1
        while rest > height:
            pages += 1
            rest -= height
        y = rest

    def new_page() -> None:
        nonlocal pages, y
        if y > 0:
            pages += 1
            y = 0.0

    def paragraph_height(p, avail: float, in_table: bool = False) -> tuple[float, float]:
        style = p.style or normal
        size, spacing, before, after = styles.metrics(style)
        pf = p.paragraph_format
        if pf.line_spacing is not None:
            spacing = pf.line_spacing if isinstance(pf.line_spacing, float) else -pf.line_spacing.pt
        if pf.space_after is not None:
            after = pf.space_after.pt
        if pf.space_before is not None:
            before = pf.space_before.pt
        extents = [e for e in p._p.findall(f".//{qn('wp:inline')}/{qn('wp:extent')}")]
        if extents:
            widths = [int(e.get("cx")) / EMU_PER_PT for e in extents]
            heights = [int(e.get("cy")) / EMU_PER_PT for e in extents]
            total = max(heights) if sum(widths) <= avail + 1 else sum(heights)
            return total + before + after, 0.0
        text_w = 0.0
        max_size = size
        for run in p.runs:
            run_size = run.font.size.pt if run.font.size is not None else size
            max_size = max(max_size, run_size) if run.text.strip() else max_size
            text_w += _text_width(run.text, run_size) * (1.06 if run.bold else 1.0)
        for link in p._p.findall(qn("w:hyperlink")):
            text_w += _text_width("".join(t.text or "" for t in link.iter(_W_T)), size)
        if p._p.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is not None:
            avail -= 36
        natural = max_size * SINGLE_LINE_FACTOR
        line_h = max(natural, -spacing) if spacing < 0 else natural * spacing
        lines = max(1, math.ceil(text_w * 1.03 / max(avail, 36)))
        return lines * line_h + before + after, line_h

    def child_height(tr) -> float:
        el = tr.find(f"{qn('w:trPr')}/{qn('w:trHeight')}")
        return int(el.get(qn("w:val"), 0)) / 20 if el is not None else 0.0

    for child in doc.element.body.iterchildren():
        if child.tag == _W_P:
            from docx.text.paragraph import Paragraph as DocxParagraph

            p = DocxParagraph(child, doc)
            if p.paragraph_format.page_break_before or child.findall(
                f".//{qn('w:br')}[@{qn('w:type')}='page']"
            ):
                new_page()
            h, line_h = paragraph_height(p, width / columns)
            place(h / columns, line_h)
            if child.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None:
                sec = min(sec + 1, len(sections) - 1)
                width, height = sections[sec]
                columns = layouts[sec][0]
                if not layouts[sec][1]:  # a continuous break stays on the page
                    pages += 1
                    y = 0.0
        elif child.tag == _W_TBL:
            from docx.table import Table as DocxTable

            table = DocxTable(child, doc)
            grid = [
                int(g.get(qn("w:w"), 0)) / 20 for g in child.findall(f"{qn('w:tblGrid')}/{qn('w:gridCol')}")
            ]
            for row in table.rows:
                row_h = 0.0
                seen = set()
                for idx, cell in enumerate(row.cells):
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    col_w = (
                        cell.width.pt
                        if cell.width is not None
                        else (grid[idx] if idx < len(grid) and grid[idx] else width / max(len(grid), 1))
                    )
                    cell_h = sum(
                        paragraph_height(cp, col_w - CELL_PADDING, True)[0] for cp in cell.paragraphs
                    )
                    row_h = max(row_h, cell_h)
                tr_h = child_height(row._tr)
                place(max(row_h, tr_h) + 1.0)
    return pages
