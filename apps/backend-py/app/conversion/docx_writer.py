"""DOCX generation with python-docx from the recovered document model."""

from __future__ import annotations

import io
import logging
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .docx_styles import ALIGN, StyleBook, apply_pitch
from .model import (
    ColumnMark,
    DocumentModel,
    HeadingStyle,
    ImageRow,
    ListInfo,
    Paragraph,
    Picture,
    Rule,
    Run,
    Table,
)

log = logging.getLogger(__name__)

_ALIGN = {k: v for k, v in ALIGN.items() if k != "left"}
EMU_PER_PT = 12700


class _Base:
    """Formatting a paragraph style already provides, so runs only carry real overrides."""

    __slots__ = ("size", "bold", "italic", "font")

    def __init__(self, size: float, bold: bool, italic: bool, font: str) -> None:
        self.size, self.bold, self.italic, self.font = size, bold, italic, font


_BULLET_GLYPHS = ("•", "◦", "▪")
_LINK_COLOR = RGBColor(0x05, 0x63, 0xC1)
# Word's "single" line spacing is roughly 1.15x the font size for common text fonts.
SINGLE_LINE_FACTOR = 1.15


def _el(tag: str, **attrs) -> OxmlElement:
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn(f"w:{k}"), str(v))
    return el


class DocxWriter:
    def __init__(self, model: DocumentModel, title: str = "") -> None:
        self.model = model
        self.doc = Document()
        self.title = title
        self._num_ids: dict[int, int] = {}
        self._avail_width = 468.0
        self._avail_height = 648.0
        self._last_was_table = False
        self._styles: StyleBook | None = None
        self._drawing_id = 1000
        self._column_break = False  # next paragraph starts the next column
        self._page_width = 468.0

    # ------------------------------------------------------------------ styles

    def _style_font(self, style, name: str, size: float, bold=None, italic=None) -> None:
        font = style.font
        font.name = name
        font.size = Pt(size)
        font.bold = bold
        font.italic = italic
        font.color.rgb = RGBColor(0, 0, 0)
        rfonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            rfonts.attrib.pop(qn(f"w:{attr}"), None)
        rfonts.set(qn("w:eastAsia"), name)
        rfonts.set(qn("w:cs"), name)

    def _setup_styles(self) -> None:
        m = self.model
        styles = self.doc.styles
        normal = styles["Normal"]
        self._style_font(normal, m.body_font, m.body_size)
        pf = normal.paragraph_format
        # Vertical spacing is reproduced per paragraph from the source geometry.
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        if m.line_pitch:
            multiple = m.line_pitch / (SINGLE_LINE_FACTOR * m.body_size)
            pf.line_spacing = round(max(0.9, min(2.0, multiple)), 2)
        else:
            pf.line_spacing = 1.0

        title = m.title_style or HeadingStyle(m.body_size * 1.7, True, False)
        self._style_font(styles["Title"], m.body_font, title.size, title.bold, title.italic)
        tpf = styles["Title"].paragraph_format
        tpf.space_before, tpf.space_after, tpf.line_spacing = Pt(0), Pt(0), 1.0
        ppr = styles["Title"].element.get_or_add_pPr()
        for border in ppr.findall(qn("w:pBdr")):
            ppr.remove(border)  # the template draws a rule under the title

        defaults = {1: 1.6, 2: 1.3, 3: 1.15}
        for level in (1, 2, 3):
            hs = m.heading_styles.get(level) or HeadingStyle(m.body_size * defaults[level], True, False)
            style = styles[f"Heading {level}"]
            self._style_font(style, m.body_font, hs.size, hs.bold, hs.italic)
            hpf = style.paragraph_format
            hpf.space_before, hpf.space_after, hpf.line_spacing = Pt(0), Pt(0), 1.0
            hpf.keep_with_next = True

        styles["List Paragraph"].paragraph_format.space_after = Pt(0)
        self._styles = StyleBook(self.doc, m, self._style_font)

    # --------------------------------------------------------------- numbering

    def _numbering_root(self):
        return self.doc.part.numbering_part.element

    def _create_list(self, info: ListInfo) -> int:
        root = self._numbering_root()
        abstract_ids = [int(a.get(qn("w:abstractNumId"))) for a in root.findall(qn("w:abstractNum"))]
        num_ids = [int(n.get(qn("w:numId"))) for n in root.findall(qn("w:num"))]
        abstract_id = max(abstract_ids, default=0) + 1
        num_id = max(num_ids, default=0) + 1

        abstract = _el("w:abstractNum", abstractNumId=abstract_id)
        abstract.append(_el("w:multiLevelType", val="hybridMultilevel"))
        for lvl in range(9):
            ordered, fmt, start, suffix = info.levels.get(
                lvl, (info.ordered, "decimal" if info.ordered else "bullet", 1, ".")
            )
            node = _el("w:lvl", ilvl=lvl)
            node.append(_el("w:start", val=start if ordered else 1))
            node.append(_el("w:numFmt", val=fmt if ordered else "bullet"))
            text = f"%{lvl + 1}{suffix}" if ordered else _BULLET_GLYPHS[lvl % 3]
            node.append(_el("w:lvlText", val=text))
            node.append(_el("w:lvlJc", val="left"))
            ppr = OxmlElement("w:pPr")
            ppr.append(_el("w:ind", left=720 + 360 * lvl, hanging=360))
            node.append(ppr)
            abstract.append(node)

        first_num = root.find(qn("w:num"))
        if first_num is not None:
            first_num.addprevious(abstract)
        else:
            root.append(abstract)
        num = _el("w:num", numId=num_id)
        num.append(_el("w:abstractNumId", val=abstract_id))
        root.append(num)
        return num_id

    # -------------------------------------------------------------------- runs

    def _add_run(self, paragraph, run: Run, base: HeadingStyle, parent=None):
        r = paragraph.add_run(run.text)
        if parent is not None:  # python-docx appends to the paragraph; move under the hyperlink
            parent.append(r._r)
        if run.bold != base.bold:
            r.bold = run.bold
        if run.italic != base.italic:
            r.italic = run.italic
        if run.underline or run.link:
            r.underline = True
        if run.size and abs(run.size - base.size) > 0.2:
            r.font.size = Pt(round(run.size * 2) / 2)
        if run.font and run.font != getattr(base, "font", self.model.body_font):
            r.font.name = run.font
        if run.link:
            r.font.color.rgb = _LINK_COLOR
        elif run.color:
            r.font.color.rgb = RGBColor((run.color >> 16) & 255, (run.color >> 8) & 255, run.color & 255)
        if run.superscript:
            r.font.superscript = True
        if run.subscript:
            r.font.subscript = True
        return r

    def _add_runs(self, paragraph, runs: list[Run], base: HeadingStyle) -> None:
        link_el = None
        link_url = None
        for run in runs:
            if run.link:
                if link_url != run.link:
                    rel = self.doc.part.relate_to(run.link, RT.HYPERLINK, is_external=True)
                    link_el = OxmlElement("w:hyperlink")
                    link_el.set(qn("r:id"), rel)
                    link_el.set(qn("w:history"), "1")
                    paragraph._p.append(link_el)
                    link_url = run.link
                self._add_run(paragraph, run, base, parent=link_el)
            else:
                link_el = link_url = None
                self._add_run(paragraph, run, base)

    # ------------------------------------------------------------------ blocks

    def _after_table(self, paragraph=None) -> None:
        """Keep content clear of a preceding table."""
        if self._last_was_table:
            if paragraph is None:  # adjacent tables would fuse into one in Word
                spacer = self.doc.add_paragraph()
                spacer.paragraph_format.space_after = Pt(0)
            elif paragraph.style.name in ("Normal", "List Paragraph"):
                paragraph.paragraph_format.space_before = Pt(6)
        self._last_was_table = False

    def _add_paragraph(self, para: Paragraph) -> None:
        m = self.model
        style_align, style_left, style_first, style_pitch = "left", 0.0, 0.0, None
        if para.kind == "title":
            p = self.doc.add_paragraph(style="Title")
            hs = m.title_style or HeadingStyle(para.size, True, False)
            base = _Base(hs.size, hs.bold, hs.italic, m.body_font)
        elif para.kind == "heading":
            p = self.doc.add_paragraph(style=f"Heading {para.level}")
            hs = m.heading_styles.get(para.level) or HeadingStyle(para.size, True, False)
            base = _Base(hs.size, hs.bold, hs.italic, m.body_font)
        elif para.kind == "list":
            p = self.doc.add_paragraph(style="List Paragraph")
            base = _Base(m.body_size, False, False, m.body_font)
            if para.list_id not in self._num_ids:
                self._num_ids[para.list_id] = self._create_list(m.lists[para.list_id])
            num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
            num_pr.get_or_add_ilvl().val = para.level
            num_pr.get_or_add_numId().val = self._num_ids[para.list_id]
        else:
            name, key = self._styles.resolve(para)
            p = self.doc.add_paragraph(style=name)
            base = _Base(key.size, key.bold, key.italic, key.font)
            style_align, style_left, style_first, style_pitch = (
                key.align, key.indent_left, key.first_indent, key.pitch,
            )  # fmt: skip
        self._after_table(p)
        pf = p.paragraph_format
        if para.align != style_align:
            p.alignment = ALIGN[para.align]
        if para.kind != "list":
            if abs(para.indent_left - style_left) >= 1:
                pf.left_indent = Pt(para.indent_left)
            if abs(para.first_indent - style_first) >= 1:
                pf.first_line_indent = Pt(para.first_indent)
        if para.line_pitch and (style_pitch is None or abs(para.line_pitch - style_pitch) > 0.6):
            if para.kind not in ("title", "heading") or len(para.runs) > 0:
                apply_pitch(pf, para.line_pitch, para.size)
        if para.space_before >= 1:
            pf.space_before = Pt(round(para.space_before * 2) / 2)
        if para.keep_with_next:
            pf.keep_with_next = True
        if para.page_break_before:
            pf.page_break_before = True
        self._start_column(p)
        for picture in para.anchors:
            self._add_anchor(p, picture)
        self._add_runs(p, para.runs, base)

    def _start_column(self, paragraph) -> None:
        if self._column_break:
            paragraph.add_run().add_break(WD_BREAK.COLUMN)
            self._column_break = False

    def _set_columns(self, docx_section, count: int, gap: float) -> None:
        sect = docx_section._sectPr
        cols = sect.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols")
            sect.append(cols)
        for attr in ("num", "space", "equalWidth"):
            cols.attrib.pop(qn(f"w:{attr}"), None)
        if count > 1:
            cols.set(qn("w:num"), str(count))
            cols.set(qn("w:space"), str(int(round(gap * 20))))
            cols.set(qn("w:equalWidth"), "1")

    def _add_column_mark(self, mark: ColumnMark) -> None:
        """Multi-column zones are continuous Word sections with real columns."""
        if mark.kind == "break":
            self._column_break = True
            return
        if self._last_was_table:
            self.doc.add_paragraph().paragraph_format.space_after = Pt(0)
            self._last_was_table = False
        new_page = mark.page_break_before
        section = self.doc.add_section(WD_SECTION.NEW_PAGE if new_page else WD_SECTION.CONTINUOUS)
        self._compact_section_break()
        if mark.kind == "start":
            gap = mark.gap
            if mark.offset > 0 and mark.count == 2:
                # equal columns whose left edges are `offset` apart, as in the source
                gap = max(6.0, min(2 * mark.offset - self._page_width, self._page_width / 2))
            self._set_columns(section, mark.count, gap)
            self._avail_width = (self._page_width - gap * (mark.count - 1)) / mark.count
        else:
            self._set_columns(section, 1, 0)
            self._avail_width = self._page_width
            self._column_break = False

    def _picture_run(self, paragraph, pic: Picture, width: float, height: float | None = None):
        """Add an inline picture run, re-encoding once if python-docx rejects the bytes."""
        data = pic.data
        for attempt in range(2):
            run = paragraph.add_run()
            try:
                run.add_picture(io.BytesIO(data), width=Pt(width), height=Pt(height) if height else None)
                return run
            except Exception:  # noqa: BLE001
                run._r.getparent().remove(run._r)
                if attempt:
                    log.warning("Could not embed image", exc_info=True)
                    return None
                try:
                    data = pymupdf.Pixmap(pic.data).tobytes("png")
                except Exception:  # noqa: BLE001
                    log.warning("Could not re-encode image", exc_info=True)
                    return None
        return None

    def _add_anchor(self, paragraph, pic: Picture) -> None:
        """Float a picture at its page coordinates (full-bleed, wrapped or margin-crossing)."""
        run = self._picture_run(paragraph, pic, pic.width, pic.height)
        if run is None:
            return
        inline = run._r.find(f".//{qn('wp:inline')}")
        self._drawing_id += 1
        anchor = OxmlElement("wp:anchor")
        for name, value in (
            ("distT", 0), ("distB", 0), ("distL", 114300), ("distR", 114300), ("simplePos", 0),
            ("relativeHeight", self._drawing_id), ("behindDoc", 1 if pic.wrap == "behind" else 0),
            ("locked", 0), ("layoutInCell", 1), ("allowOverlap", 1),
        ):  # fmt: skip
            anchor.set(name, str(value))
        simple = OxmlElement("wp:simplePos")
        simple.set("x", "0")
        simple.set("y", "0")
        anchor.append(simple)
        for tag, offset in (("wp:positionH", pic.x), ("wp:positionV", pic.y)):
            pos = OxmlElement(tag)
            pos.set("relativeFrom", "page")
            off = OxmlElement("wp:posOffset")
            off.text = str(int(round(offset * EMU_PER_PT)))
            pos.append(off)
            anchor.append(pos)
        children = {child.tag: child for child in inline}
        anchor.append(children[qn("wp:extent")])
        effect = OxmlElement("wp:effectExtent")
        for side in ("l", "t", "r", "b"):
            effect.set(side, "0")
        anchor.append(effect)
        if pic.wrap == "square":
            wrap = OxmlElement("wp:wrapSquare")
            wrap.set("wrapText", "bothSides")
        elif pic.wrap == "topbottom":
            wrap = OxmlElement("wp:wrapTopAndBottom")
        else:
            wrap = OxmlElement("wp:wrapNone")
        anchor.append(wrap)
        for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
            if qn(tag) in children:
                anchor.append(children[qn(tag)])
        inline.getparent().replace(inline, anchor)

    def _add_rule(self, rule: Rule) -> None:
        p = self.doc.add_paragraph()
        self._after_table(p)
        pf = p.paragraph_format
        pf.line_spacing = Pt(2)
        pf.space_before = Pt(max(0.0, round(rule.space_before)))
        pf.left_indent = Pt(rule.indent_left)
        right = self._avail_width - rule.indent_left - rule.width
        if right > 1:
            pf.right_indent = Pt(right)
        if rule.page_break_before:
            pf.page_break_before = True
        borders = OxmlElement("w:pBdr")
        borders.append(_el("w:bottom", val="single", sz=6, space=1, color="auto"))
        p._p.get_or_add_pPr().append(borders)

    def _add_images(self, row: ImageRow) -> None:
        p = self.doc.add_paragraph()
        self._after_table(p)
        self._start_column(p)
        pf = p.paragraph_format
        pf.line_spacing = 1.0
        pf.keep_with_next = row.keep_with_next  # keep the picture with its caption
        if row.align in _ALIGN:
            p.alignment = _ALIGN[row.align]
        elif row.indent_left >= 1:
            pf.left_indent = Pt(row.indent_left)
        if row.space_before >= 1:
            pf.space_before = Pt(round(row.space_before * 2) / 2)
        if row.page_break_before:
            pf.page_break_before = True
        gap = 6.0 * (len(row.images) - 1)
        total = sum(i.width for i in row.images) + gap
        avail = self._avail_width - (row.indent_left if row.align == "left" else 0.0)
        scale = min(1.0, avail / total) if total else 1.0
        scale = min([scale] + [0.95 * self._avail_height / i.height for i in row.images if i.height])
        for n, pic in enumerate(row.images):
            if n:
                p.add_run("  ")
            self._picture_run(p, pic, pic.width * scale, pic.height * scale)

    def _add_table(self, table: Table) -> None:
        if table.page_break_before:
            self._last_was_table = False
            self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        if self._column_break:  # a table cannot carry the break itself
            self._last_was_table = False
            holder = self.doc.add_paragraph()
            holder.paragraph_format.space_after = Pt(0)
            self._start_column(holder)
        self._after_table()
        n_rows, n_cols = len(table.rows), table.n_cols
        if not n_rows or not n_cols:
            return
        t = self.doc.add_table(rows=n_rows, cols=n_cols)
        if table.ruled:
            t.style = "Table Grid"
        t.autofit = False
        if table.align == "center":
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
        total = sum(table.col_widths) or 1.0
        scale = min(1.0, self._avail_width / total)
        widths = [Pt(w * scale) for w in table.col_widths]
        base = HeadingStyle(self.model.body_size, False, False)

        for r, (row, docx_row) in enumerate(zip(table.rows, t.rows)):
            cells = docx_row.cells
            if r < table.header_rows:
                docx_row._tr.get_or_add_trPr().append(_el("w:tblHeader"))
            if r < len(table.row_heights) and table.row_heights[r] > 0:
                docx_row.height = Pt(round(table.row_heights[r], 1))
                docx_row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            for c, cell in enumerate(row):
                cells[c].width = widths[c]
                if cell is None:
                    continue
                if table.row_heights:
                    cells[c].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = cells[c].paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                self._add_runs(p, cell.runs, base)
        for r, row in enumerate(table.rows):  # merge after filling so indices stay valid
            for c, cell in enumerate(row):
                if cell is not None and (cell.row_span > 1 or cell.col_span > 1):
                    r1 = min(r + cell.row_span, n_rows) - 1
                    c1 = min(c + cell.col_span, n_cols) - 1
                    if (r1, c1) != (r, c):
                        try:
                            merged = t.cell(r, c).merge(t.cell(r1, c1))
                        except Exception:  # noqa: BLE001 - keep the cells unmerged, keep the text
                            log.warning("Table merge skipped", extra={"cell": (r, c)})
                            continue
                        for extra in merged.paragraphs[1:]:
                            if not extra.text:
                                extra._p.getparent().remove(extra._p)
        self._last_was_table = True

    # ---------------------------------------------------------------- sections

    def _apply_section(self, docx_section, section) -> None:
        landscape = section.width > section.height
        docx_section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        docx_section.page_width = Pt(section.width)
        docx_section.page_height = Pt(section.height)
        docx_section.left_margin = Pt(section.margin_left)
        docx_section.right_margin = Pt(section.margin_right)
        docx_section.top_margin = Pt(section.margin_top)
        docx_section.bottom_margin = Pt(section.margin_bottom)
        docx_section.header_distance = Pt(min(36.0, max(18.0, section.margin_top / 2)))
        docx_section.footer_distance = Pt(min(36.0, max(18.0, section.margin_bottom / 2)))
        self._avail_width = section.width - section.margin_left - section.margin_right
        self._page_width = self._avail_width
        self._avail_height = section.height - section.margin_top - section.margin_bottom

    def _compact_section_break(self) -> None:
        """Fold python-docx's empty section-break paragraph into the paragraph before it."""
        body = self.doc.element.body
        holders = [
            p for p in body.findall(qn("w:p")) if p.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None
        ]
        if not holders:
            return
        holder = holders[-1]
        prev = holder.getprevious()
        if prev is None or prev.tag != qn("w:p") or holder.findall(qn("w:r")):
            return  # a table must be followed by a paragraph, so the holder stays
        if prev.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None:
            return
        sect = holder.find(f"{qn('w:pPr')}/{qn('w:sectPr')}")
        prev.get_or_add_pPr().append(sect)
        body.remove(holder)
        self._last_was_table = False

    def _add_furniture(self) -> None:
        m = self.model
        first = self.doc.sections[0]
        small = HeadingStyle(m.body_size, False, False)
        if m.header:
            p = first.header.paragraphs[0]
            self._add_runs(p, m.header, small)
        if m.footer or m.footer_page_field:
            p = first.footer.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if m.footer_page_field:
                prefix, suffix = m.footer_page_field
                template = m.footer[0] if m.footer else Run("")
                size = Pt(template.size) if template.size else None

                def styled(text: str):
                    run = p.add_run(text)
                    run.font.size = size
                    return run

                if prefix:
                    styled(prefix)
                field = _el("w:fldSimple", instr="PAGE")
                number = styled("1")
                field.append(number._r)
                p._p.append(field)
                if suffix:
                    styled(suffix)
            else:
                self._add_runs(p, m.footer, small)

    # -------------------------------------------------------------------- main

    def build(self) -> Document:
        self._setup_styles()
        body = self.doc.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)
        for index, section in enumerate(self.model.sections):
            if index == 0:
                docx_section = self.doc.sections[0]
            else:
                docx_section = self.doc.add_section(WD_SECTION.NEW_PAGE)
                self._compact_section_break()
            self._apply_section(docx_section, section)
            self._last_was_table = False if index == 0 else self._last_was_table
            for position, block in enumerate(section.blocks):
                following = section.blocks[position + 1] if position + 1 < len(section.blocks) else None
                if (
                    isinstance(block, ColumnMark)
                    and block.kind == "end"
                    and isinstance(following, ColumnMark)
                    and following.kind == "start"
                ):
                    # one section break, not two: an empty paragraph between back-to-back
                    # breaks would spill onto a page of its own when the page is full
                    self._column_break = False
                    continue
                if isinstance(block, Paragraph):
                    self._add_paragraph(block)
                elif isinstance(block, ImageRow):
                    self._add_images(block)
                elif isinstance(block, Table):
                    self._add_table(block)
                elif isinstance(block, Rule):
                    self._add_rule(block)
                elif isinstance(block, ColumnMark):
                    self._add_column_mark(block)
        if self._last_was_table or not self.doc.paragraphs:
            self.doc.add_paragraph()
        self._add_furniture()
        props = self.doc.core_properties
        props.title = self.title
        props.author = "DocLift"
        props.comments = "Converted from PDF by DocLift"
        return self.doc


def write_docx(model: DocumentModel, out_path: Path, title: str = "") -> None:
    DocxWriter(model, title).build().save(str(out_path))
