"""Fidelity report for one conversion.

A conversion is not "successful" because the DOCX opens. The report compares the source PDF
with the generated DOCX (and, when a renderer is available, with the DOCX rendered back to
PDF) page by page, and lists concrete problems.
"""

from __future__ import annotations

import json
import re
from collections import Counter
from dataclasses import asdict, dataclass, field
from pathlib import Path

import pymupdf
from docx import Document
from docx.oxml.ns import qn

from ..conversion import convert_pdf_to_docx
from ..conversion.extract import MIN_IMAGE_SIDE, extract_page
from ..conversion.inspect_docx import _paragraph_text, inspect_docx
from ..conversion.text_repair import FontCatalog, count_suspicious, find_suspicious, repair_pages
from .render import RendererUnavailable, docx_to_pdf, find_renderer

EMU_PER_PT = 12700
GRID = (36, 46)  # columns, rows of the ink-density grid used for visual comparison


@dataclass(slots=True)
class PageComparison:
    page: int
    text_match: float  # share of the source page's words found on the same output page
    visual_similarity: float | None  # weighted Jaccard of ink density, 0..1


@dataclass(slots=True)
class QAReport:
    source: str
    layout_mode: str
    source_pages: int
    output_pages: int  # rendered page count, or the estimate when no renderer is available
    rendered: bool
    explicit_page_breaks: int
    text_coverage: float  # share of source letters/digits present in the output
    suspicious_in_source: int  # raw extraction, before repair
    suspicious_in_output: int
    suspicious_samples: list[str]
    glyphs_repaired: dict[str, str]
    images_in_source: int
    images_in_output: int
    floating_images: int
    image_size_mismatches: int
    headings: dict[str, int]
    styles_used: dict[str, int]
    header_text: str
    footer_text: str
    pages: list[PageComparison] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.problems

    def to_json(self) -> str:
        return json.dumps(asdict(self), ensure_ascii=False, indent=2)


def _alnum(text: str) -> Counter:
    return Counter(ch.lower() for ch in text if ch.isalnum())


def _coverage(source: str, output: str) -> float:
    a, b = _alnum(source), _alnum(output)
    total = sum(a.values())
    return 1.0 if not total else sum(min(n, b[ch]) for ch, n in a.items()) / total


_TOKEN = re.compile(r"[^\W_]+")


def _words(text: str) -> Counter:
    # letter/digit runs, so leader dots or punctuation glued to a word do not hide it
    return Counter(w.lower() for w in _TOKEN.findall(text))


def _page_text_match(source: str, output: str) -> float:
    a, b = _words(source), _words(output)
    total = sum(a.values())
    return 1.0 if not total else sum(min(n, b[w]) for w, n in a.items()) / total


def ink_grid(page: pymupdf.Page, grid: tuple[int, int] = GRID) -> list[int]:
    """Darkness per cell (0 = blank, 255 = solid) of a page rendered straight to grid size."""
    cols, rows = grid
    matrix = pymupdf.Matrix(cols / page.rect.width, rows / page.rect.height)
    pix = page.get_pixmap(matrix=matrix, colorspace=pymupdf.csGRAY, alpha=False)
    return [255 - v for v in pix.samples[: pix.width * pix.height]]


def visual_similarity(a: pymupdf.Page, b: pymupdf.Page) -> float:
    """Weighted Jaccard similarity of where the ink is on two pages (1.0 = same layout)."""
    ga, gb = ink_grid(a), ink_grid(b)
    n = min(len(ga), len(gb))
    num = sum(min(ga[i], gb[i]) for i in range(n))
    den = sum(max(ga[i], gb[i]) for i in range(n))
    return 1.0 if den == 0 else num / den


def side_by_side(a: pymupdf.Page, b: pymupdf.Page, out: Path, dpi: int = 60) -> Path:
    """PNG with the source page on the left and the rendered output on the right."""
    sheet = pymupdf.open()
    gap = 12.0
    width = a.rect.width + b.rect.width + gap
    height = max(a.rect.height, b.rect.height)
    page = sheet.new_page(width=width, height=height)
    page.insert_image(pymupdf.Rect(0, 0, a.rect.width, a.rect.height), pixmap=a.get_pixmap(dpi=dpi))
    page.insert_image(pymupdf.Rect(a.rect.width + gap, 0, width, b.rect.height), pixmap=b.get_pixmap(dpi=dpi))
    page.get_pixmap(dpi=dpi).save(out)
    sheet.close()
    return out


def representative_pages(count: int, limit: int = 12) -> list[int]:
    """First pages (cover, front matter, contents), evenly spaced samples, and the last pages."""
    if count <= limit:
        return list(range(count))
    picks = set(range(min(4, count))) | {count - 2, count - 1}
    step = (count - 1) / (limit - len(picks) + 1)
    picks |= {round(step * i) for i in range(1, limit - len(picks) + 1)}
    return sorted(p for p in picks if 0 <= p < count)[:limit]


def _docx_facts(docx_path: Path) -> dict:
    doc = Document(str(docx_path))
    body = doc.element.body
    paragraphs = list(body.iter(qn("w:p")))
    text = "\n".join(_paragraph_text(p) for p in paragraphs)
    section = doc.sections[0]
    header = " ".join(p.text for p in section.header.paragraphs).strip()
    footer = " ".join(_paragraph_text(p._p) for p in section.footer.paragraphs).strip()
    sizes = []
    for drawing in body.iter(qn("wp:inline"), qn("wp:anchor")):
        ext = drawing.find(qn("wp:extent"))
        sizes.append((int(ext.get("cx")) / EMU_PER_PT, int(ext.get("cy")) / EMU_PER_PT))
    styles = Counter()
    for p in paragraphs:
        style = p.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
        styles[style.get(qn("w:val")) if style is not None else "Normal"] += 1
    return {
        "text": text + "\n" + header + "\n" + footer,
        "header": header,
        "footer": footer,
        "image_sizes": sizes,
        "floating": len(list(body.iter(qn("wp:anchor")))),
        "styles": dict(styles.most_common()),
    }


def _source_facts(pdf_path: Path) -> dict:
    with pymupdf.open(str(pdf_path)) as pdf:
        raw_text = "\n".join(page.get_text() for page in pdf)
        catalog = FontCatalog(pdf)
        pages = [extract_page(pdf[i], i, catalog) for i in range(pdf.page_count)]
        repair_pages(pdf, pages)
        page_text = []
        for page in pages:
            parts = [ln.text for ln in page.lines]
            parts += [
                "".join(s.text for s in spans) for t in page.tables for c in t.cells for spans in c.spans
            ]
            page_text.append(" ".join(parts))
        sizes = []
        for page in pdf:
            for info in page.get_image_info():
                x0, y0, x1, y1 = info["bbox"]
                if x1 - x0 >= MIN_IMAGE_SIDE and y1 - y0 >= MIN_IMAGE_SIDE:
                    sizes.append((x1 - x0, y1 - y0))
        return {
            "pages": pdf.page_count,
            "raw_text": raw_text,
            "page_text": page_text,
            "image_sizes": sizes,
        }


def _size_mismatches(source: list[tuple[float, float]], output: list[tuple[float, float]]) -> int:
    """Pictures whose displayed size or aspect ratio changed (order-independent match)."""
    remaining = list(output)
    missing = 0
    for w, h in source:
        hit = next(
            (
                o
                for o in remaining
                if abs(o[0] - w) <= max(2.0, 0.02 * w) and abs(o[1] - h) <= max(2.0, 0.02 * h)
            ),
            None,
        )
        if hit is None:  # a uniformly scaled picture keeps its aspect ratio but not its size
            missing += 1
        else:
            remaining.remove(hit)
    return missing


def analyse(
    pdf_path: Path,
    docx_path: Path,
    rendered_pdf: Path | None = None,
    quality: dict | None = None,
    compare_dir: Path | None = None,
    visual_pages: int = 12,
) -> QAReport:
    src = _source_facts(pdf_path)
    out = _docx_facts(docx_path)
    stats = inspect_docx(docx_path)
    quality = quality or {}
    report = QAReport(
        source=pdf_path.name,
        layout_mode=quality.get("layout_mode", "unknown"),
        source_pages=src["pages"],
        # without a render: explicit breaks in faithful output, simulated layout otherwise
        output_pages=stats.explicit_pages if quality.get("layout_mode") == "faithful" else stats.page_count,
        rendered=False,
        explicit_page_breaks=stats.explicit_pages - 1,
        text_coverage=round(_coverage(" ".join(src["page_text"]), out["text"]), 4),
        suspicious_in_source=count_suspicious(src["raw_text"]),
        suspicious_in_output=count_suspicious(out["text"]),
        suspicious_samples=find_suspicious(out["text"], 8),
        glyphs_repaired=quality.get("glyphs_repaired", {}),
        images_in_source=len(src["image_sizes"]),
        images_in_output=len(out["image_sizes"]),
        floating_images=out["floating"],
        image_size_mismatches=_size_mismatches(src["image_sizes"], out["image_sizes"]),
        headings={f"h{k}": v for k, v in sorted(stats.headings.items())},
        styles_used=out["styles"],
        header_text=out["header"],
        footer_text=out["footer"],
    )

    if rendered_pdf is not None:
        with pymupdf.open(str(pdf_path)) as a, pymupdf.open(str(rendered_pdf)) as b:
            report.rendered = True
            report.output_pages = b.page_count
            visual = set(representative_pages(min(a.page_count, b.page_count), visual_pages))
            for i in range(min(a.page_count, b.page_count)):
                sim = None
                if i in visual:
                    sim = round(visual_similarity(a[i], b[i]), 3)
                    if compare_dir is not None:
                        compare_dir.mkdir(parents=True, exist_ok=True)
                        side_by_side(a[i], b[i], compare_dir / f"page_{i + 1:04d}.png")
                match = round(_page_text_match(src["page_text"][i], b[i].get_text()), 3)
                report.pages.append(PageComparison(i + 1, match, sim))

    _judge(report)
    return report


def _judge(r: QAReport) -> None:
    p = r.problems
    faithful = r.layout_mode == "faithful"
    if r.suspicious_in_output:
        p.append(
            f"{r.suspicious_in_output} suspicious characters in the output, e.g. {r.suspicious_samples[:3]}"
        )
    if r.text_coverage < 0.97:
        p.append(f"text coverage {r.text_coverage:.1%}: source text is missing from the output")
    if r.images_in_output != r.images_in_source:
        p.append(f"image count {r.images_in_output} != {r.images_in_source} in the source")
    if r.image_size_mismatches:
        p.append(f"{r.image_size_mismatches} images changed size or aspect ratio")
    if faithful and r.source_pages > 1 and r.explicit_page_breaks < r.source_pages - 1:
        p.append(f"only {r.explicit_page_breaks} page breaks for {r.source_pages} source pages")
    if r.rendered:
        if faithful and r.output_pages != r.source_pages:
            p.append(f"renders to {r.output_pages} pages, source has {r.source_pages}")
        drifted = [c.page for c in r.pages if c.text_match < 0.9]
        if faithful and drifted:
            p.append(f"text is on a different page than in the source on pages {drifted[:15]}")
        blurry = [c.page for c in r.pages if c.visual_similarity is not None and c.visual_similarity < 0.35]
        if faithful and blurry:
            p.append(f"layout differs visibly from the source on pages {blurry[:15]}")


def run_qa(
    pdf_path: Path,
    out_dir: Path,
    layout_mode: str = "faithful",
    render: bool = True,
    visual_pages: int = 12,
) -> QAReport:
    """Convert, render (when possible), analyse; artefacts are written to `out_dir`."""
    out_dir.mkdir(parents=True, exist_ok=True)
    docx_path = out_dir / (pdf_path.stem + ".docx")
    result = convert_pdf_to_docx(pdf_path, docx_path, layout_mode=layout_mode)
    rendered = None
    if render and find_renderer() is not None:
        try:
            rendered = docx_to_pdf(docx_path, out_dir / "render")
        except RendererUnavailable:
            rendered = None
    report = analyse(pdf_path, docx_path, rendered, result.quality, out_dir / "compare", visual_pages)
    (out_dir / "report.json").write_text(report.to_json(), encoding="utf-8")
    return report
