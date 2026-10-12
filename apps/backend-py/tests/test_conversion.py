"""Unit tests for extraction, reading order, structure recovery and table detection."""

from __future__ import annotations

from pathlib import Path

import pymupdf
import pytest
from docx import Document
from docx.oxml.ns import qn

from app.conversion import PHASES, convert_pdf_to_docx
from app.conversion.extract import (
    Char,
    Line,
    Span,
    count_words,
    extract_page,
    table_from_grid,
    word_font_name,
)
from app.conversion.inspect_docx import inspect_docx
from app.conversion.layout import find_gutter, split_leaves
from app.conversion.model import ImageRow, Paragraph, Run, Table
from app.conversion.structure import Row, detect_marker, normalise_runs, recover_structure
from app.conversion.text_repair import FontCatalog, repair_pages

from .conftest import sample


def load(name: str, mode: str = "flow"):
    """Extract, repair and structure a sample; returns (raw pages, document model).

    The structural tests below describe the reflowable ("flow") reading of the samples;
    layout-faithful behaviour is covered in test_layout.py.
    """
    with pymupdf.open(sample(name)) as pdf:
        catalog = FontCatalog(pdf)
        pages = [extract_page(pdf[i], i, catalog) for i in range(pdf.page_count)]
        repair_pages(pdf, pages)
        return pages, recover_structure(pages, pdf, layout_mode=mode)


def blocks(model):
    return list(model.iter_blocks())


def paragraphs(model, kind=None):
    return [b for b in blocks(model) if isinstance(b, Paragraph) and (kind is None or b.kind == kind)]


def tables(model):
    return [b for b in blocks(model) if isinstance(b, Table)]


def grid(table: Table):
    return [[c.text if c is not None else None for c in row] for row in table.rows]


def make_line(text, x0, y0, *, size=10.0, bold=False, block=0, width=None) -> Line:
    width = width if width is not None else len(text) * size * 0.5
    step = width / max(len(text), 1)
    chars = [Char(c, x0 + i * step, x0 + (i + 1) * step) for i, c in enumerate(text)]
    span = Span(chars, size, "Helvetica-Bold" if bold else "Helvetica", bold, False, False, None, y0 + size)
    return Line([span], (x0, y0, x0 + width, y0 + size * 1.2), block, 0)


# ------------------------------------------------------------------- extraction


def test_extract_character_formatting_and_links():
    pages, _ = load("03_")
    line = next(ln for ln in pages[0].lines if ln.text.startswith("Text can be"))
    by_text = {s.text.strip(" ,."): s for s in line.spans}
    assert by_text["bold"].bold and not by_text["bold"].italic
    assert by_text["italic"].italic and not by_text["italic"].bold
    assert by_text["bold italic"].bold and by_text["bold italic"].italic
    underlined = [s.text for s in line.spans if s.underline]
    assert underlined == ["underlined"]  # the drawn rule covers exactly this word
    link = [s for s in line.spans if s.link]
    assert [s.text for s in link] == ["example.com"] and link[0].link == "https://example.com"


def test_extract_images_with_positions():
    pages, _ = load("05_")
    assert [len(p.images) for p in pages] == [3, 1]
    banner = pages[0].images[0]
    assert banner.ext == "png" and banner.data[:8] == b"\x89PNG\r\n\x1a\n"
    assert banner.bbox[2] - banner.bbox[0] == pytest.approx(432, abs=1)


def test_extract_ruled_tables_removes_cell_text_from_flow():
    pages, _ = load("04_")
    assert [(t.n_rows, t.n_cols) for t in pages[0].tables] == [(5, 4), (5, 4), (11, 5)]
    flow_text = " ".join(ln.text for ln in pages[0].lines)
    assert "Tomato" not in flow_text and "Rainy season" not in flow_text
    assert "Name:" in flow_text  # the borderless table is not ruled; it stays in the flow for now


def test_table_from_grid_recovers_spans():
    boxes = [(0, 0, 10, 10), (10, 0, 30, 10), (0, 10, 10, 20), (10, 10, 20, 20), (20, 10, 30, 20)]
    table = table_from_grid(boxes, (0, 0, 30, 20), 0, ruled=True)
    assert (table.n_rows, table.n_cols) == (2, 3)
    wide = next(c for c in table.cells if (c.row, c.col) == (0, 1))
    assert (wide.row_span, wide.col_span) == (1, 2)


def test_word_count_ignores_symbol_tokens():
    assert count_words(["•", "Hello", "—", "x2", "…", "GH₵", "250"]) == 4


@pytest.mark.parametrize(
    ("pdf_font", "expected"),
    [
        ("Helvetica-BoldOblique", "Arial"),
        ("ABCDEF+TimesNewRomanPS-BoldMT", "Times New Roman"),
        ("DejaVuSans", "DejaVu Sans"),
        ("Courier", "Courier New"),
        ("Lato-Bold", "Lato"),
    ],
)
def test_word_font_name(pdf_font, expected):
    assert word_font_name(pdf_font) == expected


# ---------------------------------------------------------------- reading order


def test_two_columns_are_read_left_then_right():
    prose = "lorem ipsum dolor sit amet consectetur"
    left = [make_line(prose, 60, 100 + 14 * i, block=1, width=220) for i in range(8)]
    right = [make_line(prose, 320, 100 + 14 * i, block=2, width=220) for i in range(8)]
    title = make_line("A Centred Title Across Both Columns", 150, 60, size=18, block=0, width=320)
    interleaved = [title] + [ln for pair in zip(left, right) for ln in pair]

    gutter = find_gutter(interleaved)
    assert gutter is not None and 280 <= gutter[0] and gutter[1] <= 320
    leaves, multi = split_leaves(interleaved)
    assert multi
    assert leaves == [[title], left, right]


def test_single_column_and_list_markers_are_not_columns():
    body = [
        make_line("a full width line of ordinary body text here", 60, 100 + 14 * i, block=1, width=480)
        for i in range(10)
    ]
    assert split_leaves(body) == ([body], False)
    # bullets with hanging text look like two x-clusters but must not become columns
    items = []
    for i in range(8):
        items.append(make_line("•", 60, 100 + 14 * i, block=i, width=4))
        items.append(make_line("list item text", 78, 100 + 14 * i, block=i, width=90))
    assert split_leaves(items)[1] is False


def test_two_column_sample_reading_order():
    _, model = load("02_")
    assert model.multi_column_pages == 2
    texts = [p.text for p in paragraphs(model)]
    assert texts[0] == "Freshness Signals in Online Produce Markets"
    assert texts[1] == "A. Mensah and K. Owusu"
    headings = [p.text for p in paragraphs(model, "heading")]
    assert headings == ["Section 1", "Section 2", "Section 3", "Section 4"]
    # no paragraph was cut at a column or page boundary, and none was interleaved
    body = [p.text for p in paragraphs(model, "body")][1:]
    assert all(t.startswith("Agriculture remains") and t.endswith("with customers.") for t in body)
    assert len(body) == 14


def test_mixed_layout_reading_order():
    _, model = load("06_")
    kinds = []
    for b in blocks(model):
        if isinstance(b, Paragraph):
            kinds.append(b.text[:12] if b.kind in ("title", "heading") else b.kind)
        else:
            kinds.append(type(b).__name__)
    assert kinds == [
        "Quarterly Bu",
        "body",
        "Highlights",
        "body",
        "list",
        "list",
        "list",
        "ImageRow",
        "caption",
        "Pricing",
        "body",
        "Table",
        "Outlook",
        "body",
    ]
    assert model.multi_column_pages == 1


# ---------------------------------------------------------- structure recovery


def test_headings_follow_font_size_tiers():
    _, model = load("03_")
    title = paragraphs(model, "title")
    assert [p.text for p in title] == ["Project Handbook"] and title[0].align == "center"
    levels = {p.text: p.level for p in paragraphs(model, "heading")}
    assert levels == {
        "1. Overview": 1,
        "2. Process": 1,
        "3. Styles": 1,
        "1.1 Goals": 2,
        "2.1 Nested example": 2,
        "1.1.1 Non-goals": 3,
    }
    assert model.heading_counts == {1: 3, 2: 2, 3: 1}
    assert (model.body_size, model.body_font) == (10.5, "Arial")
    assert [model.heading_styles[i].size for i in (1, 2, 3)] == [18.0, 14.0, 12.0]


def test_lists_are_detected_with_nesting_and_numbering():
    _, model = load("03_")
    items = paragraphs(model, "list")
    assert (model.ordered_lists, model.unordered_lists) == (1, 3)
    ordered = [p for p in items if p.ordered]
    assert [p.text for p in ordered] == [
        "Upload the PDF",
        "Wait for conversion",
        "Download the .docx",
        "Open and review in Word",
    ]
    info = model.lists[ordered[0].list_id]
    assert info.ordered and info.levels[0][:3] == (True, "decimal", 1)
    nested = [(p.text, p.level) for p in items if p.list_id == items[-1].list_id]
    assert nested == [
        ("Frontend", 0),
        ("Next.js upload form", 1),
        ("Progress indicator", 1),
        ("Backend", 0),
        ("Conversion endpoint", 1),
        ("Download endpoint", 1),
    ]
    assert not any(p.text.startswith(("•", "1 ")) for p in items)  # glyphs are not kept as text


@pytest.mark.parametrize(
    ("segments", "expected"),
    [
        (["•", "Bullet text"], (False, True)),
        (["1", "Bare number"], (True, False)),
        (["2.", "Dotted number"], (True, True)),
        (["(b)", "Lettered"], (True, True)),
        (["• inline bullet"], (False, True)),
        (["- dash item"], (False, False)),
        (["3) inline number"], (True, False)),
        (["Plain sentence here"], None),
        (["A. Mensah and K. Owusu"], (True, False)),  # weak: needs a sibling to count
        (["2024 was a good year"], None),
        (["Name:", "Nana"], None),
    ],
)
def test_detect_marker(segments, expected):
    x = 60.0
    lines = []
    for text in segments:
        line = make_line(text, x, 100)
        lines.append(line)
        x = line.x1 + 12
    marker = detect_marker(Row(lines), body_size=10.0)
    assert (None if marker is None else (marker.ordered, marker.strong)) == expected


def test_heading_sized_numbered_line_is_not_a_list_item():
    row = Row([make_line("1. Overview", 60, 100, size=18, bold=True)])
    assert detect_marker(row, body_size=10.5) is None


def test_paragraphs_reflow_and_continue_across_pages():
    _, model = load("01_")
    body = paragraphs(model, "body")
    assert len(body) == 8
    assert all(p.text.endswith("with customers.") for p in body)
    assert all("\n" not in p.text for p in body)
    assert len(model.sections) == 1


def test_inline_formatting_becomes_runs():
    _, model = load("03_")
    para = next(p for p in paragraphs(model) if p.text.startswith("Text can be"))
    fmt = {r.text.strip(" ,."): (r.bold, r.italic, r.underline, r.link) for r in para.runs}
    assert fmt["bold"] == (True, False, False, None)
    assert fmt["italic"] == (False, True, False, None)
    assert fmt["underlined"] == (False, False, True, None)
    assert fmt["bold italic"] == (True, True, False, None)
    assert fmt["example.com"][3] == "https://example.com"


def test_special_characters_and_scripts():
    _, model = load("08_")
    text = "\n".join(p.text for p in paragraphs(model))
    for needle in (
        "café, naïve, señor, Müller, Ångström",
        "Ɛyɛ adwuma, ɔkɔ fie",
        "GH₵ 250, € 20, £ 15",
        "— – “quotes” ‘single’ …",
        "≤ ≥ ≠ ± √ ∑ π",
        "office, affluent, floor",
    ):
        assert needle in text
    math = next(p for p in paragraphs(model) if p.text.startswith("Math markup"))
    assert math.text.startswith("Math markup: H2O, CO2, x2 + y2 = z2")
    assert [r.text for r in math.runs if r.subscript] == ["2", "2"]
    assert [r.text for r in math.runs if r.superscript] == ["2", "2", "2"]


def test_header_footer_and_footnotes():
    _, model = load("12_")
    assert "".join(r.text for r in model.header) == "ACME Research - Confidential Draft"
    assert model.footer_page_field == ("Page ", "")
    body_text = [p.text for p in paragraphs(model)]
    assert not any("ACME Research" in t or t == "Page 1" for t in body_text)
    notes = [p for p in paragraphs(model) if "Footnote-style note" in p.text]
    assert [p.text[:24] for p in notes] == [f"{n} Footnote-style note {n}:" for n in range(1, 5)]
    assert all(p.runs[0].superscript for p in notes)


def test_page_sizes_and_orientation_become_sections():
    _, model = load("11_")
    assert [(round(s.width), round(s.height)) for s in model.sections] == [(612, 792), (842, 595), (612, 792)]
    assert [len(tables(model)), tables(model)[0].n_cols, len(tables(model)[0].rows)] == [1, 8, 8]
    assert tables(model)[0] in model.sections[1].blocks


def test_short_pages_keep_their_page_breaks():
    _, model = load("10_")
    chapters = paragraphs(model, "heading")
    assert [p.text for p in chapters] == [f"Chapter {n}" for n in range(1, 41)]
    assert [p.page_break_before for p in chapters] == [False] + [True] * 39


def test_images_stay_in_flow_order():
    _, model = load("05_")
    sequence = []
    for b in blocks(model):
        if isinstance(b, ImageRow):
            sequence.append(f"<{len(b.images)} image>")
        elif b.kind == "heading" or b.text.startswith(("Figure", "Text")):
            sequence.append(b.text[:8])
    assert sequence == [
        "Full-wid",
        "<1 image>",
        "Figure 1",
        "Two imag",
        "<2 image>",
        "Figure 2",
        "High-res",
        "<1 image>",
        "Text fol",
    ]
    assert model.images_detected == 4
    square = next(b for b in blocks(model) if isinstance(b, ImageRow) and len(b.images) == 2).images[0]
    assert square.width == pytest.approx(square.height, abs=0.5)


def test_normalise_runs_merges_and_strips_control_characters():
    runs = normalise_runs([Run("Hello "), Run(" wor\x00ld", bold=False), Run("  "), Run("!", bold=True)])
    assert [(r.text, r.bold) for r in runs] == [("Hello world ", False), ("!", True)]


# ------------------------------------------------------------- table detection


def test_ruled_tables_keep_cells_and_header():
    _, model = load("04_")
    a = tables(model)[0]
    assert grid(a) == [
        ["Crop", "Region", "Price (GHS/kg)", "Freshness"],
        ["Tomato", "Ashanti", "12.50", "High"],
        ["Maize", "Northern", "4.20", "Medium"],
        ["Plantain", "Western", "6.80", "High"],
        ["Yam", "Bono East", "9.00", "Medium"],
    ]
    assert a.ruled and a.header_rows == 1
    assert all(r.bold for c in a.rows[0] for r in c.runs)


def test_merged_cells_are_preserved():
    _, model = load("04_")
    b = tables(model)[1]
    # "Sales" spans two columns; "Quarter" and "Notes" span the two header rows
    assert grid(b)[:3] == [
        ["Quarter", "Sales", None, "Notes"],
        [None, "Online", "In-person", None],
        ["Q1", "120", "300", "Rainy season"],
    ]
    assert [(c.row_span, c.col_span) if c else None for c in b.rows[0]] == [(2, 1), (1, 2), None, (2, 1)]


def test_borderless_table_is_detected():
    _, model = load("04_")
    c = tables(model)[2]
    assert not c.ruled
    assert grid(c) == [["Name:", "Nana"], ["Role:", "Frontend"]]


def test_table_continued_across_pages_is_one_table():
    _, model = load("04_")
    assert model.tables_detected == 4
    d = tables(model)[3]
    rows = grid(d)
    assert rows[0] == ["#", "Item", "Qty", "Unit cost", "Total"] and d.header_rows == 1
    assert [r[0] for r in rows[1:]] == [str(n) for n in range(1, 61)]  # none lost, none duplicated
    assert rows[60] == ["60", "Item 60", "5", "90.00", "450.00"]


@pytest.mark.parametrize("name", ["01_", "02_", "03_", "05_", "08_", "10_", "12_"])
def test_prose_is_not_mistaken_for_tables(name):
    _, model = load(name)
    assert model.tables_detected == 0


# -------------------------------------------------------------- synthetic PDFs


def build_pdf(path: Path, lines: list[tuple[float, float, str, float]]) -> Path:
    doc = pymupdf.open()
    page = doc.new_page()
    for x, y, text, size in lines:
        page.insert_text((x, y), text, fontsize=size, fontname="helv")
    doc.save(path)
    doc.close()
    return path


def test_inline_markers_and_weak_marker_validation(tmp_path):
    lines = [(72, 80, "Intro paragraph before the list.", 11)]
    lines += [
        (72, 110 + 16 * i, t, 11)
        for i, t in enumerate(["3. third thing", "4. fourth thing", "5. fifth thing"])
    ]
    lines += [(72, 180, "A. Mensah wrote this sentence.", 11)]
    lines += [(72, 210 + 16 * i, t, 11) for i, t in enumerate(["- dash one", "- dash two"])]
    lines += [(72, 255, "A closing remark ends the list.", 11)]
    lines += [(72, 285, "- a lone dash is just punctuation", 11)]
    with pymupdf.open(build_pdf(tmp_path / "lists.pdf", lines)) as pdf:
        model = recover_structure([extract_page(pdf[0], 0)], pdf, layout_mode="flow")
    items = paragraphs(model, "list")
    assert [p.text for p in items] == ["third thing", "fourth thing", "fifth thing", "dash one", "dash two"]
    assert (model.ordered_lists, model.unordered_lists) == (1, 1)
    assert model.lists[items[0].list_id].levels[0][:3] == (True, "decimal", 3)  # numbering starts at 3
    body = [p.text for p in paragraphs(model, "body")]
    assert "A. Mensah wrote this sentence." in body
    assert "- a lone dash is just punctuation" in body


def test_whitespace_aligned_table_from_synthetic_pdf(tmp_path):
    lines = [(72, 80, "Quarterly figures follow in a small layout table.", 11)]
    for i, row in enumerate(
        [("Region", "Units", "Revenue"), ("North", "120", "4,300"), ("South", "95", "3,150")]
    ):
        for x, text in zip((72, 220, 360), row):
            lines.append((x, 120 + 18 * i, text, 11))
    with pymupdf.open(build_pdf(tmp_path / "table.pdf", lines)) as pdf:
        model = recover_structure([extract_page(pdf[0], 0)], pdf, layout_mode="flow")
    assert model.tables_detected == 1
    assert grid(tables(model)[0]) == [
        ["Region", "Units", "Revenue"],
        ["North", "120", "4,300"],
        ["South", "95", "3,150"],
    ]


# ------------------------------------------------------- end-to-end conversion

EXPECTED_METRICS = {
    "01_": dict(
        source_word_count=838,
        output_word_count=838,
        headings_detected=0,
        tables_detected=0,
        images_detected=0,
        multi_column_page_count=0,
        output_page_count=2,
    ),
    "02_": dict(
        source_word_count=747,
        output_word_count=747,
        headings_detected=4,
        h1_heading_count=4,
        multi_column_page_count=2,
        tables_detected=0,
    ),
    "03_": dict(
        source_word_count=123,
        output_word_count=123,
        ordered_lists_detected=1,
        unordered_lists_detected=3,
        ordered_lists_reconstructed=1,
        unordered_lists_reconstructed=3,
        headings_detected=6,
        h1_heading_count=3,
        h2_heading_count=2,
        h3_heading_count=1,
        output_page_count=1,
    ),
    # 10 fewer output words: the header row repeated on pages 2 and 3 is emitted once
    "04_": dict(
        source_word_count=441,
        output_word_count=431,
        tables_detected=4,
        headings_detected=4,
        output_page_count=3,
    ),
    "05_": dict(source_word_count=138, output_word_count=138, images_detected=4, headings_detected=3),
    "06_": dict(
        source_word_count=239,
        output_word_count=239,
        unordered_lists_detected=1,
        unordered_lists_reconstructed=1,
        tables_detected=1,
        images_detected=1,
        headings_detected=3,
        multi_column_page_count=1,
        output_page_count=1,
    ),
    "08_": dict(source_word_count=47, output_word_count=47, headings_detected=0, output_page_count=1),
    "10_": dict(
        source_word_count=14434,
        output_word_count=14434,
        headings_detected=40,
        h1_heading_count=40,
        output_page_count=40,
    ),
    "11_": dict(source_word_count=279, output_word_count=279, tables_detected=1, output_page_count=3),
    "12_": dict(source_word_count=462, output_word_count=462, headings_detected=0),
}


@pytest.mark.parametrize("name", sorted(EXPECTED_METRICS))
def test_metrics_per_sample(name, tmp_path):
    out = tmp_path / "out.docx"
    result = convert_pdf_to_docx(sample(name), out, layout_mode="flow")
    assert out.stat().st_size == result.output_size > 0
    assert set(result.metrics) == {
        "source_word_count",
        "output_word_count",
        "ordered_lists_detected",
        "unordered_lists_detected",
        "ordered_lists_reconstructed",
        "unordered_lists_reconstructed",
        "headings_detected",
        "h1_heading_count",
        "h2_heading_count",
        "h3_heading_count",
        "tables_detected",
        "images_detected",
        "multi_column_page_count",
        "output_page_count",
    }
    for key, value in EXPECTED_METRICS[name].items():
        assert result.metrics[key] == value, key
    m = result.metrics
    assert m["headings_detected"] == m["h1_heading_count"] + m["h2_heading_count"] + m["h3_heading_count"]
    # what the report says was detected is what the generated file really contains
    stats = inspect_docx(out)
    assert stats.tables == m["tables_detected"]
    assert stats.images == m["images_detected"]
    assert sum(stats.headings.values()) == m["headings_detected"]


def test_progress_phases_run_in_order(tmp_path):
    seen: list[tuple[str, int]] = []
    convert_pdf_to_docx(sample("10_"), tmp_path / "o.docx", lambda phase, pct: seen.append((phase, pct)))
    order = [name for name, _, _ in PHASES]
    phases = list(dict.fromkeys(p for p, _ in seen))
    assert phases == order[1:-1]  # QUEUED and COMPLETED are set by the job service
    percents = [pct for _, pct in seen]
    assert percents == sorted(percents) and percents[0] > 0 and percents[-1] < 100


def test_generated_docx_uses_real_word_structures(tmp_path):
    out = tmp_path / "o.docx"
    convert_pdf_to_docx(sample("03_"), out, layout_mode="flow")
    doc = Document(str(out))
    styles = [p.style.name for p in doc.paragraphs]
    assert styles[0] == "Title"
    assert [styles.count(f"Heading {n}") for n in (1, 2, 3)] == [3, 2, 1]
    numbered = [p for p in doc.paragraphs if p._p.find(f"{qn('w:pPr')}/{qn('w:numPr')}") is not None]
    assert len(numbered) == 15 and all(p.style.name == "List Paragraph" for p in numbered)
    levels = [p._p.find(f".//{qn('w:ilvl')}").get(qn("w:val")) for p in numbered]
    assert levels[-6:] == ["0", "1", "1", "0", "1", "1"]
    styled = next(p for p in doc.paragraphs if p.text.startswith("Text can be"))
    assert [r.text for r in styled.runs if r.bold and not r.italic] == ["bold"]
    assert [r.text for r in styled.runs if r.underline] == ["underlined"]
    link = styled._p.find(qn("w:hyperlink"))
    rel = doc.part.rels[link.get(qn("r:id"))]
    assert rel.is_external and rel.target_ref == "https://example.com"
    assert doc.sections[0].page_width.pt == pytest.approx(595.3, abs=0.5)  # A4 preserved


def test_generated_docx_tables_sections_and_furniture(tmp_path):
    out = tmp_path / "t.docx"
    convert_pdf_to_docx(sample("04_"), out, layout_mode="flow")
    doc = Document(str(out))
    assert [(len(t.rows), len(t.columns)) for t in doc.tables] == [(5, 4), (5, 4), (2, 2), (61, 5)]
    merged = doc.tables[1].rows[0].cells
    assert merged[1]._tc is merged[2]._tc and merged[1].text == "Sales"
    assert doc.tables[3]._tbl.findall(f".//{qn('w:tblHeader')}")  # header repeats on each page

    convert_pdf_to_docx(sample("11_"), out, layout_mode="flow")
    doc = Document(str(out))
    assert [s.orientation for s in doc.sections] == [0, 1, 0]
    assert len(doc.tables[0].columns) == 8

    convert_pdf_to_docx(sample("12_"), out, layout_mode="flow")
    doc = Document(str(out))
    section = doc.sections[0]
    assert section.header.paragraphs[0].text == "ACME Research - Confidential Draft"
    footer = section.footer.paragraphs[0]
    assert footer._p.find(qn("w:fldSimple")).get(qn("w:instr")) == "PAGE"
    assert "ACME Research" not in "\n".join(p.text for p in doc.paragraphs)


def test_scanned_page_is_embedded_not_dropped(tmp_path):
    # Uploads of fully scanned PDFs are rejected up front; the converter itself stays safe.
    result = convert_pdf_to_docx(sample("07_"), tmp_path / "s.docx", layout_mode="flow")
    assert result.metrics["images_detected"] == 1 and result.metrics["source_word_count"] == 0


def test_large_document_converts_quickly(tmp_path):
    result = convert_pdf_to_docx(sample("10_"), tmp_path / "big.docx", layout_mode="flow")
    assert result.page_count == 40
    assert result.seconds < 30
